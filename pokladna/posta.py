"""Odesílání e-mailů přes SMTP: uzávěrka (souhrn + měsíční Excel) a faktury (PDF).

Automatická uzávěrka: když je vyplněné `uzaverka_auto_cas`, odejde každý den v ten čas
(nebo po zapnutí serveru, pokud se v ten čas nestihla) na adresy z `uzaverka_prijemci`.
"""

from __future__ import annotations

import html
import io
import json
import logging
import smtplib
import ssl
import threading
from datetime import datetime
from email.message import EmailMessage
from email.utils import formataddr, make_msgid
from typing import List, Optional

from . import db, uzaverka
from .eet_klient import PRAHA
from .fakturace import kc

log = logging.getLogger("pokladna.posta")
STAV_SOUBOR = db.DB_CESTA.parent / "posta_stav.json"
MESICE = ["leden", "únor", "březen", "duben", "květen", "červen", "červenec", "srpen", "září", "říjen", "listopad", "prosinec"]


class ChybaPosty(Exception):
    pass


def adresy(text: str) -> List[str]:
    return [a.strip() for a in (text or "").replace(";", ",").split(",") if "@" in a]


def odesli(nast: dict, komu: List[str], predmet: str, html_telo: str, prilohy: Optional[list] = None) -> None:
    """prilohy = [(nazev, bytes, mime)]"""
    server = (nast.get("smtp_server") or "").strip()
    if not server:
        raise ChybaPosty("Není nastavený SMTP server (Nastavení → E-mail)")
    if not komu:
        raise ChybaPosty("Chybí adresa příjemce")
    odesilatel = (nast.get("smtp_odesilatel") or nast.get("smtp_uzivatel") or "").strip()
    if "@" not in odesilatel:
        raise ChybaPosty("Vyplňte adresu odesílatele (Nastavení → E-mail)")
    msg = EmailMessage()
    msg["Subject"] = predmet
    msg["From"] = formataddr((nast.get("firma_nazev") or "Pokladna", odesilatel))
    msg["To"] = ", ".join(komu)
    msg["Message-ID"] = make_msgid(domain=odesilatel.split("@")[-1])
    msg.set_content("Tento e-mail je ve formátu HTML.")
    msg.add_alternative(html_telo, subtype="html")
    for nazev, data, mime in prilohy or []:
        hlavni, vedlejsi = mime.split("/", 1)
        msg.add_attachment(data, maintype=hlavni, subtype=vedlejsi, filename=nazev)

    port = int(nast.get("smtp_port") or 587)
    zab = (nast.get("smtp_zabezpeceni") or "STARTTLS").upper()
    try:
        if zab == "SSL":
            spojeni = smtplib.SMTP_SSL(server, port, timeout=30, context=ssl.create_default_context())
        else:
            spojeni = smtplib.SMTP(server, port, timeout=30)
        with spojeni as s:
            if zab == "STARTTLS":
                s.starttls(context=ssl.create_default_context())
            if nast.get("smtp_uzivatel"):
                s.login(nast["smtp_uzivatel"], nast.get("smtp_heslo") or "")
            s.send_message(msg)
    except smtplib.SMTPAuthenticationError:
        raise ChybaPosty("SMTP server odmítl přihlášení – zkontrolujte jméno a heslo")
    except (smtplib.SMTPException, OSError) as e:
        raise ChybaPosty(f"E-mail se nepodařilo odeslat: {e}")


# --------------------------------------------------------------------------- uzávěrka

def _datum(iso: str) -> str:
    d = datetime.fromisoformat(iso)
    return f"{d.day}. {d.month}. {d.year}"


def uzaverka_html(s: dict, nast: dict, uctenky: list) -> str:
    h = s["platby"].get("HOTOVE") or {}
    k = s["platby"].get("KARTA") or {}
    st = s["typy"].get("STORNO") or {}
    pr = s["typy"].get("PRODEJ") or {}

    def bunka(nazev, hodnota):
        return (f'<td style="padding:10px 14px;background:#f4f4f6;border-radius:8px"><div style="color:#5b5b62;font-size:12px">{nazev}</div>'
                f'<div style="font-size:20px;font-weight:bold">{hodnota}</div></td>')

    dph = "".join(f"<tr><td>{r['sazba_dph']} %</td><td align=right>{kc(r['zaklad'])}</td><td align=right>{kc(r['dph'])}</td>"
                  f"<td align=right>{kc(r['celkem'])}</td></tr>" for r in s["dph"])
    prodavaci = "".join(f"<tr><td>{html.escape(p['jmeno'] or '')}</td><td align=right>{p['pocet']}</td><td align=right>{kc(p['celkem'])}</td></tr>"
                        for p in s["prodavaci"])
    faktury = "".join(f"<tr><td>{({'PREVOD': 'převodem', 'DOBIRKA': 'dobírkou', 'HOTOVE': 'hotově', 'KARTA': 'kartou'})[f['forma_uhrady']]}</td>"
                      f"<td align=right>{f['pocet']}</td><td align=right>{kc(f['celkem'])}</td></tr>" for f in s["faktury"])
    radky = "".join(f"<tr><td>{u['cas'][11:16]}</td><td>{html.escape(u['cislo'])}{' (storno)' if u['typ'] == 'STORNO' else ''}</td>"
                    f"<td>{html.escape(u['prodavac'] or '')}</td><td>{'hotově' if u['platba'] == 'HOTOVE' else 'kartou'}</td>"
                    f"<td align=right>{kc(u['celkem'])}</td></tr>" for u in uctenky)
    eet = ", ".join(f"{k2} {v}" for k2, v in s["eet"].items()) or "—"
    tab = 'style="border-collapse:collapse;width:100%;font-size:13px" cellpadding=5'
    return f"""<div style="font-family:Arial,sans-serif;color:#1d1d1f;max-width:680px">
<h2 style="margin:0 0 4px">Uzávěrka {_datum(s['datum'])}</h2>
<div style="color:#5b5b62;margin-bottom:14px">{html.escape(nast.get('firma_nazev') or '')} · {html.escape(nast.get('provozovna') or '')}</div>
<table style="border-spacing:8px;margin-left:-8px"><tr>{bunka('Tržba celkem', kc(s['celkem']))}{bunka('Hotově', kc(h.get('celkem') or 0))}{bunka('Kartou', kc(k.get('celkem') or 0))}</tr>
<tr>{bunka('Prodejek', pr.get('pocet', 0))}{bunka('Storna', f"{st.get('pocet', 0)} · {kc(st.get('celkem') or 0)}")}{bunka('Zaokrouhlení', kc(s['zaokrouhleni']))}</tr></table>
{f'<h3>DPH</h3><table {tab}><tr><th align=left>Sazba</th><th align=right>Základ</th><th align=right>DPH</th><th align=right>Celkem</th></tr>{dph}</table>' if nast.get('platce_dph') == '1' and dph else ''}
{f'<h3>Prodavači</h3><table {tab}>{prodavaci}</table>' if prodavaci else ''}
{f'<h3>Vystavené faktury</h3><table {tab}>{faktury}</table>' if faktury else ''}
<h3>Prodejky</h3><table {tab}>{radky or '<tr><td>Žádné prodejky</td></tr>'}</table>
<p style="color:#777;font-size:12px">EET: {html.escape(eet)} · Měsíční přehled je v příloze (Excel).</p></div>"""


def mesicni_xlsx(con, datum: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    d = datetime.fromisoformat(datum)
    ws.title = f"{MESICE[d.month - 1]} {d.year}"
    hlavicka = ["Datum", "Prodejek", "Storen", "Hotově", "Kartou", "Celkem", "Základ 21 %", "DPH 21 %", "Základ 12 %", "DPH 12 %",
                "Bez DPH (0 %)", "Faktury převodem"]
    ws.append(hlavicka)
    dny = uzaverka.mesic(con, datum)
    for r in dny:
        ws.append([datetime.fromisoformat(r["datum"]).date(), r["uctenek"], r["storen"]] +
                  [r[k] / 100 for k in ("hotove", "kartou", "celkem", "zaklad21", "dph21", "zaklad12", "dph12", "zaklad0", "faktury_prevodem")])
    ws.append(["Celkem", sum(r["uctenek"] for r in dny), sum(r["storen"] for r in dny)] +
              [sum(r[k] for r in dny) / 100 for k in ("hotove", "kartou", "celkem", "zaklad21", "dph21", "zaklad12", "dph12", "zaklad0", "faktury_prevodem")])
    cervena = PatternFill("solid", fgColor="D9302A")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cervena
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for radek in ws.iter_rows(min_row=2):
        radek[0].number_format = "D.M.YYYY"
        for c in radek[3:]:
            c.number_format = '#,##0.00 "Kč"'
    for i, w in enumerate([12, 10, 8, 14, 14, 14, 14, 12, 14, 12, 14, 16], 1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.freeze_panes = "A2"

    u = wb.create_sheet(f"Prodejky {d.day}.{d.month}.")
    u.append(["Čas", "Číslo", "Typ", "Prodavač", "Platba", "Celkem", "EET"])
    for r in uzaverka.uctenky(con, datum):
        u.append([r["cas"][11:16], r["cislo"], "storno" if r["typ"] == "STORNO" else "prodej", r["prodavac"],
                  "hotově" if r["platba"] == "HOTOVE" else "kartou", r["celkem"] / 100, r["eet_stav"]])
    z = wb.create_sheet(f"Zboží {d.day}.{d.month}.")
    z.append(["Zboží", "EAN", "Množství", "Tržba s DPH"])
    for r in uzaverka.prodane_zbozi(con, datum):
        z.append([r["nazev"], r["ean"], r["mnozstvi"], r["celkem"] / 100])
    for list_ in (u, z):
        for c in list_[1]:
            c.font = Font(bold=True)
        list_.column_dimensions["A"].width = 38 if list_ is z else 8
        list_.column_dimensions["B"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def odesli_uzaverku(datum: Optional[str] = None, komu: Optional[List[str]] = None) -> dict:
    with db.spojeni() as con:
        nast = db.nastaveni(con)
        datum = datum or db.ted()[:10]
        prijemci = komu or adresy(nast.get("uzaverka_prijemci") or "")
        if not prijemci:
            raise ChybaPosty("Vyplňte příjemce uzávěrky (Nastavení → E-mail)")
        s = uzaverka.souhrn(con, datum)
        telo = uzaverka_html(s, nast, uzaverka.uctenky(con, datum))
        xlsx = mesicni_xlsx(con, datum)
    d = datetime.fromisoformat(datum)
    misto = nast.get("provozovna") or nast.get("firma_nazev") or "pokladna"
    odesli(nast, prijemci, f"Uzávěrka {misto} {d.day}. {d.month}. {d.year}: {kc(s['celkem'])}", telo,
           [(f"report {MESICE[d.month - 1]} {d.year}.xlsx", xlsx,
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")])
    stav = _cti_stav()
    stav["uzaverka"] = {"datum": datum, "odeslano": db.ted(), "komu": prijemci}
    _uloz_stav(stav)
    log.info("Uzávěrka %s odeslána na %s", datum, ", ".join(prijemci))
    return stav["uzaverka"]


def _cti_stav() -> dict:
    try:
        return json.loads(STAV_SOUBOR.read_text())
    except (OSError, ValueError):
        return {}


def _uloz_stav(stav: dict) -> None:
    STAV_SOUBOR.write_text(json.dumps(stav, ensure_ascii=False, indent=1))


def stav() -> dict:
    return _cti_stav()


class Planovac:
    def __init__(self) -> None:
        self._stop = threading.Event()
        self._vlakno: Optional[threading.Thread] = None

    def _smycka(self) -> None:
        while not self._stop.wait(60):
            try:
                with db.spojeni() as con:
                    nast = db.nastaveni(con)
                cas = (nast.get("uzaverka_auto_cas") or "").strip()
                if not cas or not adresy(nast.get("uzaverka_prijemci") or ""):
                    continue
                ted = datetime.now(PRAHA)
                hh, mm = cas.split(":")
                if (ted.hour, ted.minute) < (int(hh), int(mm)):
                    continue
                dnes = ted.date().isoformat()
                if (_cti_stav().get("uzaverka") or {}).get("datum") == dnes:
                    continue
                odesli_uzaverku(dnes)
            except ChybaPosty as e:
                log.warning("Automatická uzávěrka: %s", e)
                self._stop.wait(1800)          # při chybě nezkoušet každou minutu
            except Exception:
                log.exception("Automatická uzávěrka selhala")

    def spust(self) -> None:
        if self._vlakno is None:
            self._vlakno = threading.Thread(target=self._smycka, name="posta", daemon=True)
            self._vlakno.start()

    def zastav(self) -> None:
        self._stop.set()


planovac = Planovac()
