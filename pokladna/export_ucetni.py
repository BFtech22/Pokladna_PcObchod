"""Měsíční podklady pro účetní v jednom ZIPu.

- přehled.xlsx: souhrn DPH, tržby po dnech, prodejky, vydané doklady, příjemky, výdejky, pokladní doklady, stav skladu,
- vydané/: každý vydaný doklad v PDF a ISDOC (Pohoda, Money, iDoklad, ABRA ho načtou),
- přijaté/: faktury dodavatelů nahrané při naskladnění.
"""

from __future__ import annotations

import calendar
import html
import io
import re
import zipfile
from datetime import date, timedelta
from pathlib import Path
from typing import Callable

from . import db, fakturace, kasa, uzaverka
from .fakturace import kc

NAZVY_TYPU = {"FAKTURA": "faktura", "ZALOHOVA": "zálohová faktura", "DDZP": "daňový doklad k platbě", "OPRAVNY": "opravný doklad"}
SAZBY = (21, 12, 0)


def _rozpad(polozky: list, pdp: bool = False) -> dict:
    """{sazba: (základ, DPH)} z položek dokladu."""
    return {r["sazba"]: (r["zaklad"], r["dph"]) for r in fakturace.rekap_dph(polozky, pdp)}


def _posledni_den(mesic: str) -> str:
    r, m = int(mesic[:4]), int(mesic[5:7])
    return f"{mesic}-{calendar.monthrange(r, m)[1]:02d}"


def data_mesice(con, mesic: str) -> dict:
    vzor = f"{mesic}%"
    prodejky = []
    for p in con.execute("SELECT * FROM prodeje WHERE cas LIKE ? ORDER BY id", (vzor,)).fetchall():
        pol = [dict(x) for x in con.execute("SELECT sazba_dph, celkem FROM prodej_polozky WHERE prodej_id = ?", (p["id"],))]
        prodejky.append({**dict(p), "rozpad": _rozpad(pol)})
    doklady = [r["id"] for r in con.execute(
        "SELECT id FROM faktury_vydane WHERE typ != 'NABIDKA' AND datum_vystaveni LIKE ? ORDER BY id", (vzor,))]
    prijemky = [dict(r) for r in con.execute(
        "SELECT p.*, s.nazev AS soubor_nazev, s.cesta AS soubor_cesta FROM prijemky p LEFT JOIN faktury_soubory s ON s.id = p.soubor_id "
        "WHERE p.cas LIKE ? ORDER BY p.id", (vzor,))]
    vydejky = [dict(r) for r in con.execute("SELECT * FROM vydejky WHERE cas LIKE ? ORDER BY id", (vzor,))]
    pokladni = [dict(r) for r in con.execute("SELECT * FROM pokladni_doklady WHERE cas LIKE ? ORDER BY id", (vzor,))]
    return {"prodejky": prodejky, "doklady": doklady, "prijemky": prijemky, "vydejky": vydejky, "pokladni": pokladni}


def stav_skladu(con, k_datu: str) -> list:
    """Množství na skladě ke konci dne `k_datu` (součet pohybů) oceněné poslední nákupní cenou."""
    out = []
    for r in con.execute(
            "SELECT z.id, z.nazev, z.kod, z.jednotka, z.nakupni_cena, COALESCE(SUM(p.mnozstvi), 0) AS stav FROM zbozi z "
            "LEFT JOIN pohyby p ON p.zbozi_id = z.id AND substr(p.cas, 1, 10) <= ? GROUP BY z.id ORDER BY z.nazev COLLATE NOCASE",
            (k_datu,)):
        if abs(r["stav"]) > 1e-9:
            out.append({**dict(r), "hodnota": int(round((r["nakupni_cena"] or 0) * r["stav"]))})
    return out


def zip_mesic(con, mesic: str, nast: dict, detail: Callable[[int], dict]) -> tuple:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    d = data_mesice(con, mesic)
    doklady = [detail(fid) for fid in d["doklady"]]
    konec = min(_posledni_den(mesic), db.ted()[:10])
    sklad = stav_skladu(con, konec)
    dalsi_mesic = (date.fromisoformat(_posledni_den(mesic)) + timedelta(days=1)).isoformat()
    kniha = kasa.stav(con, f"{mesic}-01"), kasa.stav(con, dalsi_mesic)

    wb = Workbook()
    cervena = PatternFill("solid", fgColor="D9302A")

    def list_(nazev, hlavicka, radky, sirky, penize_od=None):
        ws = wb.create_sheet(nazev)
        ws.append(hlavicka)
        for r in radky:
            ws.append(r)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = cervena
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, w in enumerate(sirky, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = w
        if penize_od is not None:
            for radek in ws.iter_rows(min_row=2):
                for c in radek[penize_od:]:
                    if isinstance(c.value, (int, float)):
                        c.number_format = '#,##0.00'
        ws.freeze_panes = "A2"
        return ws

    # Souhrn DPH
    souhrn = wb.active
    souhrn.title = "Souhrn"
    souhrn.append([f"{nast.get('firma_nazev') or ''} · IČO {nast.get('firma_ico') or ''} · podklady za {mesic[5:]}/{mesic[:4]}"])
    souhrn["A1"].font = Font(bold=True, size=13)
    souhrn.append([])
    souhrn.append(["Doklady", "Počet"] + [f"{k} {s} %" for s in SAZBY for k in ("Základ", "DPH")] + ["Celkem s DPH"])
    for c in souhrn[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = cervena
        c.alignment = Alignment(horizontal="center", wrap_text=True)

    def radek_souhrnu(nazev, seznam, rozpad, celkem):
        sumy = {s: [0, 0] for s in SAZBY}
        for x in seznam:
            for s, (zaklad, dph) in rozpad(x).items():
                sumy.setdefault(s, [0, 0])
                sumy[s][0] += zaklad
                sumy[s][1] += dph
        souhrn.append([nazev, len(seznam)] + [v / 100 for s in SAZBY for v in sumy[s]] + [sum(celkem(x) for x in seznam) / 100])

    radek_souhrnu("Prodejky hotově", [p for p in d["prodejky"] if p["platba"] == "HOTOVE"], lambda p: p["rozpad"], lambda p: p["celkem"])
    radek_souhrnu("Prodejky kartou", [p for p in d["prodejky"] if p["platba"] == "KARTA"], lambda p: p["rozpad"], lambda p: p["celkem"])
    popisky = {"FAKTURA": "Faktury převodem a dobírkou", "DDZP": "Daňové doklady k přijaté platbě", "OPRAVNY": "Opravné doklady",
               "ZALOHOVA": "Zálohové faktury (nejsou daňový doklad)"}
    for typ, popisek in popisky.items():
        seznam = [f for f in doklady if f["typ"] == typ and f["stav"] != "STORNO" and not f["prodej_id"]]
        radek_souhrnu(popisek, seznam, (lambda f: {}) if typ == "ZALOHOVA" else (lambda f: _rozpad(f["polozky"], bool(f["pdp"]))),
                      lambda f: f["celkem"])
    souhrn.append([])
    souhrn.append(["Faktury placené na pokladně jsou už v prodejkách (stejná tržba, EET)."])
    souhrn.append(["Přenesená daňová povinnost: základ je u faktur v řádku sazby, DPH 0 – daň přiznává odběratel."])
    souhrn.append([])
    souhrn.append(["Pokladna", "Částka"])
    souhrn.append(["Hotovost na začátku měsíce", kniha[0] / 100])
    souhrn.append(["Hotovost na konci měsíce", kniha[1] / 100])
    souhrn.append(["Hodnota skladu ke dni " + konec + " (nákupní ceny)", sum(r["hodnota"] for r in sklad) / 100])
    souhrn.column_dimensions["A"].width = 46
    for col in "BCDEFGHI":
        souhrn.column_dimensions[col].width = 14
    for radek in souhrn.iter_rows(min_row=4):
        for c in radek[1:]:
            if isinstance(c.value, float):
                c.number_format = '#,##0.00'

    dny = uzaverka.mesic(con, konec) if konec[:7] == mesic else []
    list_("Tržby po dnech", ["Datum", "Prodejek", "Storen", "Hotově", "Kartou", "Celkem", "Základ 21 %", "DPH 21 %", "Základ 12 %",
                             "DPH 12 %", "Bez DPH (0 %)", "Faktury převodem"],
          [[r["datum"], r["uctenek"], r["storen"]] + [r[k] / 100 for k in ("hotove", "kartou", "celkem", "zaklad21", "dph21", "zaklad12",
                                                                             "dph12", "zaklad0", "faktury_prevodem")] for r in dny],
          [12, 9, 8] + [13] * 9, 3)
    list_("Prodejky", ["Číslo", "Datum a čas", "Typ", "Platba", "Základ 21 %", "DPH 21 %", "Základ 12 %", "DPH 12 %", "Bez DPH",
                       "Zaokrouhlení", "Celkem", "EET", "FIK / BKP"],
          [[p["cislo"], p["cas"][:16].replace("T", " "), "storno" if p["typ"] == "STORNO" else "prodej",
            "hotově" if p["platba"] == "HOTOVE" else "kartou"]
           + [p["rozpad"].get(s, (0, 0))[i] / 100 for s in (21, 12) for i in (0, 1)] + [p["rozpad"].get(0, (0, 0))[0] / 100,
                                                                                         p["zaokrouhleni"] / 100, p["celkem"] / 100,
                                                                                         p["eet_stav"], p["eet_pok"] or ""]
           for p in d["prodejky"]],
          [14, 17, 8, 8] + [12] * 7 + [10, 30], 4)
    list_("Vydané doklady", ["Číslo", "Druh", "Vystaveno", "DUZP", "Splatnost", "Odběratel", "IČO", "DIČ", "Úhrada", "Základ 21 %",
                             "DPH 21 %", "Základ 12 %", "DPH 12 %", "Bez DPH", "Celkem", "Stav", "Uhrazeno", "PDP", "Navazuje na"],
          [[f["cislo"], NAZVY_TYPU.get(f["typ"], f["typ"]), f["datum_vystaveni"], f["duzp"] if f["typ"] != "ZALOHOVA" else "",
            f["splatnost"], f["odb_nazev"], f["odb_ico"] or "", f["odb_dic"] or "", fakturace.FORMY.get(f["forma_uhrady"], "")]
           + [(_rozpad(f["polozky"], bool(f["pdp"])).get(s, (0, 0))[i] / 100 if f["typ"] != "ZALOHOVA" else 0) for s in (21, 12) for i in (0, 1)]
           + [(_rozpad(f["polozky"], bool(f["pdp"])).get(0, (0, 0))[0] / 100 if f["typ"] != "ZALOHOVA" else 0), f["celkem"] / 100,
              {"NEUHRAZENO": "neuhrazeno", "UHRAZENO": "uhrazeno", "STORNO": "stornováno"}[f["stav"]], f["datum_uhrady"] or "",
              "ano" if f["pdp"] else "", (f["vazba"] or {}).get("cislo") or (f["prodejka"] and f"prodejka {f['prodejka']}") or ""]
           for f in doklady],
          [13, 18, 11, 11, 11, 28, 10, 13, 10] + [12] * 6 + [11, 11, 5, 16], 9)
    list_("Příjemky", ["Číslo", "Datum", "Dodavatel", "IČO dodavatele", "Faktura dodavatele", "Celkem bez DPH", "Stornováno", "Soubor faktury"],
          [[p["cislo"], p["cas"][:10], p["dodavatel"] or "", p["dodavatel_ico"] or "", p["doklad"] or "", p["celkem_bez_dph"] / 100,
            "ano" if p["stornovano"] else "", p["soubor_nazev"] or ""] for p in d["prijemky"]],
          [13, 11, 30, 12, 18, 14, 10, 30], 5)
    list_("Výdejky", ["Číslo", "Datum", "Důvod", "Komu", "Poznámka", "Hodnota v nákupních cenách", "Stornováno"],
          [[v["cislo"], v["cas"][:10], v["duvod"], v["komu"] or "", v["poznamka"] or "", v["celkem_nakup"] / 100,
            "ano" if v["stornovano"] else ""] for v in d["vydejky"]],
          [13, 11, 12, 24, 30, 16, 10], 5)
    list_("Pokladní doklady", ["Číslo", "Datum", "Příjem / výdej", "Druh", "Účel", "Osoba", "Částka", "Stornováno"],
          [[p["cislo"], p["cas"][:10], "příjem" if p["typ"] == "PRIJEM" else "výdej", kasa.DRUHY.get(p["druh"], (None, p["druh"]))[1],
            p["ucel"] or "", p["osoba"] or "", p["castka"] / 100, "ano" if p["stornovano"] else ""] for p in d["pokladni"]],
          [13, 11, 10, 18, 34, 18, 12, 10], 6)
    list_(f"Sklad k {konec[8:10]}.{konec[5:7]}.", ["Zboží", "Kód", "Množství", "MJ", "Nákupní cena bez DPH", "Hodnota"],
          [[r["nazev"], r["kod"] or "", r["stav"], r["jednotka"], (r["nakupni_cena"] or 0) / 100, r["hodnota"] / 100] for r in sklad],
          [40, 14, 10, 6, 16, 14], 4)
    xlsx = io.BytesIO()
    wb.save(xlsx)

    firma = re.sub(r"[^A-Za-z0-9]+", "_", db.bez_diakritiky(nast.get("firma_nazev") or "pokladna")).strip("_") or "pokladna"
    nazev = f"{firma}_ucetnictvi_{mesic}.zip"
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(f"prehled_{mesic}.xlsx", xlsx.getvalue())
        for f in doklady:
            try:
                z.writestr(f"vydane/{f['cislo']}.pdf", fakturace.pdf(f, nast))
            except fakturace.ChybaFakturace:
                pass
            z.writestr(f"vydane/{f['cislo']}.isdoc", fakturace.isdoc(f, nast))
        pouzite = set()
        for p in d["prijemky"]:
            if p["soubor_cesta"] and Path(p["soubor_cesta"]).is_file():
                jmeno = re.sub(r"[^\w.\- ]+", "_", f"{p['cislo']}_{p['soubor_nazev'] or 'faktura'}")
                if jmeno not in pouzite:
                    pouzite.add(jmeno)
                    z.write(p["soubor_cesta"], f"prijate/{jmeno}")
        z.writestr("CTI_ME.txt", (
            f"Podklady pro účetnictví – {nast.get('firma_nazev') or ''}, IČO {nast.get('firma_ico') or ''}, období {mesic}\n\n"
            f"prehled_{mesic}.xlsx – souhrn DPH, tržby po dnech, prodejky, vydané doklady, příjemky, výdejky, pokladní doklady a stav skladu\n"
            "vydane/ – vydané faktury, zálohové faktury, daňové doklady k platbě a opravné doklady (PDF + ISDOC pro import)\n"
            "prijate/ – faktury dodavatelů nahrané při naskladnění zboží\n\n"
            f"Vytvořeno {db.ted()[:16].replace('T', ' ')} v pokladně PcObchod.\n").encode("utf-8"))
    return buf.getvalue(), nazev


def email_html(con, mesic: str, nast: dict) -> str:
    d = data_mesice(con, mesic)
    trzby = sum(p["celkem"] for p in d["prodejky"])
    return f"""<div style="font-family:Arial,sans-serif;font-size:14px;color:#1d1d1f;max-width:620px">
<p>Dobrý den,</p>
<p>v příloze posíláme podklady pro účetnictví za <b>{mesic[5:]}/{mesic[:4]}</b>: přehled v Excelu, vydané doklady v PDF a ISDOC
a přijaté faktury dodavatelů.</p>
<ul><li>Prodejky: {len(d['prodejky'])}, tržby {kc(trzby)}</li><li>Vydané doklady: {len(d['doklady'])}</li>
<li>Příjemky: {len(d['prijemky'])}</li><li>Pokladní doklady: {len(d['pokladni'])}</li></ul>
<p>Děkujeme<br>{html.escape(nast.get('firma_nazev') or '')}</p></div>"""
